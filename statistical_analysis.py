import pandas as pd
import numpy as np
import scipy.stats as stats
import statsmodels.api as sm
from docx import Document
import os

# --- 1. GENERATE MOCK DATA ---
np.random.seed(42)
data = pd.DataFrame({
    'Sales': np.random.uniform(100, 1000, 100),
    'Discount_Applied': np.random.choice([0, 1], 100), # 0 = No, 1 = Yes
    'Profit': np.random.uniform(10, 300, 100)
})
# Adjust profit based on discount to create a real statistical effect
data.loc[data['Discount_Applied'] == 1, 'Profit'] -= np.random.uniform(20, 50, sum(data['Discount_Applied'] == 1))

# --- 2. STATISTICAL TESTS ---

# Test A: Independent T-Test (Comparing Profits with vs. without discounts)
group_no_discount = data[data['Discount_Applied'] == 0]['Profit']
group_discount = data[data['Discount_Applied'] == 1]['Profit']
t_stat, p_value_t = stats.ttest_ind(group_no_discount, group_discount)

# Test B: Pearson Correlation (Sales vs. Profit)
corr_coeff, p_value_corr = stats.pearsonr(data['Sales'], data['Profit'])

# Test C: Linear Regression (Predicting Profit using Sales)
X = sm.add_constant(data['Sales']) # Add intercept
Y = data['Profit']
model = sm.OLS(Y, X).fit()
r_squared = model.rsquared
p_value_reg = model.pvalues['Sales']

# --- 3. BUILD WORD DOCUMENT ---
doc = Document()
doc.add_heading('Statistical Analysis and Hypothesis Testing', 0)
doc.add_paragraph('Author: Faizhan Ali Syed')

# T-Test Section
doc.add_heading('1. Independent T-Test: Impact of Discounts on Profit', level=1)
doc.add_paragraph("Null Hypothesis (H0): There is no significant difference in profit between discounted and non-discounted items.")
doc.add_paragraph("Alternative Hypothesis (H1): There is a significant difference in profit.")
doc.add_paragraph(f"T-Statistic: {t_stat:.2f} | P-Value: {p_value_t:.4f}")
doc.add_paragraph("Interpretation: " + ("The p-value is < 0.05. We reject the null hypothesis; discounts significantly affect profit." if p_value_t < 0.05 else "The p-value is > 0.05. We fail to reject the null hypothesis."))

# Correlation Section
doc.add_heading('2. Pearson Correlation: Sales vs. Profit', level=1)
doc.add_paragraph("Null Hypothesis (H0): There is no linear correlation between Sales and Profit.")
doc.add_paragraph("Alternative Hypothesis (H1): There is a significant linear correlation.")
doc.add_paragraph(f"Correlation Coefficient (r): {corr_coeff:.2f} | P-Value: {p_value_corr:.4f}")
doc.add_paragraph("Interpretation: " + ("The p-value is < 0.05. We reject the null hypothesis, confirming a significant correlation." if p_value_corr < 0.05 else "No significant correlation was found."))

# Regression Section
doc.add_heading('3. Linear Regression: Predicting Profit', level=1)
doc.add_paragraph("Null Hypothesis (H0): Sales volume does not significantly predict Profit.")
doc.add_paragraph("Alternative Hypothesis (H1): Sales volume significantly predicts Profit.")
doc.add_paragraph(f"R-Squared: {r_squared:.4f} | P-Value for Sales: {p_value_reg:.4f}")
doc.add_paragraph("Interpretation: The R-squared value indicates how much variance in Profit is explained by Sales. " + ("The relationship is statistically significant." if p_value_reg < 0.05 else "The relationship is not statistically significant."))

# --- 4. SAVE DIRECTLY TO DESKTOP FOLDER ---
output_path = r"C:\Users\VICTUS\Desktop\New folder\Statistical_Analysis_Report.docx"
doc.save(output_path)
print(f"Done! Report saved to {output_path}")